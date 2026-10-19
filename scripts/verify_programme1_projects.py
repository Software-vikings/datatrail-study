#!/usr/bin/env python3
"""Execute all four Programme 1 portfolio projects in isolated temporary folders."""

from __future__ import annotations

import importlib.util
import json
import sqlite3
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PROJECTS = ROOT / "projects" / "programme-1"


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise AssertionError(f"Cannot load project module: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def verify_project_1(base: Path) -> None:
    folder = PROJECTS / "01-sales-data-quality"
    module = load_module("sales_quality_project", folder / "src" / "clean_sales.py")
    result = module.run(folder / "data" / "raw_sales.csv", base / "project-1")
    quality = result["quality"]
    assert quality["raw_rows"] == quality["accepted_rows"] + quality["rejected_rows"]
    assert quality["duplicate_transaction_ids"] == 1
    assert quality["accepted_rows"] == 8
    assert quality["rejected_rows"] == 7
    assert quality["reconciliation_passed"] is True
    assert abs(float(result["cleaned"]["revenue"].sum()) - quality["accepted_revenue"]) < 0.01
    for name in ("cleaned_sales.csv", "rejected_sales.csv", "decision_log.csv", "monthly_reconciliation.csv", "data_quality_report.json"):
        assert (base / "project-1" / name).is_file(), name
    decision = pd.read_csv(base / "project-1" / "decision_log.csv")
    assert len(decision) == quality["raw_rows"] and set(decision["decision"]) == {"accepted", "rejected"}
    source = pd.read_csv(folder / "data" / "raw_sales.csv")
    extra = pd.DataFrame([
        {"transaction_id": "T-MISSING-CUSTOMER", "order_date": "2026-03-01", "customer_id": None, "city": "Chennai", "product": "Keyboard", "quantity": 1, "unit_price": 100},
        {"transaction_id": "T-MISSING-PRODUCT", "order_date": "2026-03-02", "customer_id": "C999", "city": "Chennai", "product": None, "quantity": 1, "unit_price": 100},
        {"transaction_id": "T-INFINITE-PRICE", "order_date": "2026-03-03", "customer_id": "C998", "city": "Chennai", "product": "Mouse", "quantity": 1, "unit_price": float("inf")},
        {"transaction_id": "T-OVERFLOW-REVENUE", "order_date": "2026-03-04", "customer_id": "C997", "city": "Chennai", "product": "Monitor", "quantity": 1e308, "unit_price": 1e308},
    ])
    probe_input = base / "project-1-missing-fields.csv"
    pd.concat([source, extra], ignore_index=True).to_csv(probe_input, index=False)
    probe_result = module.run(probe_input, base / "project-1-probe")
    assert probe_result["quality"]["raw_rows"] == 19
    assert probe_result["quality"]["accepted_rows"] == 8
    assert probe_result["quality"]["rejected_rows"] == 11
    assert probe_result["quality"]["invalid_unit_prices"] == 2
    assert probe_result["quality"]["non_finite_revenue_rows"] == 1
    assert probe_result["quality"]["missing_customer_ids"] == 1
    assert probe_result["quality"]["missing_products"] == 1
    probe_decisions = pd.read_csv(base / "project-1-probe" / "decision_log.csv").set_index("transaction_id")
    assert "missing_customer_id" in probe_decisions.loc["T-MISSING-CUSTOMER", "reason"]
    assert "missing_product" in probe_decisions.loc["T-MISSING-PRODUCT", "reason"]
    assert "invalid_unit_price" in probe_decisions.loc["T-INFINITE-PRICE", "reason"]
    assert "non_finite_revenue" in probe_decisions.loc["T-OVERFLOW-REVENUE", "reason"]
    empty_sales = base / "project-1-empty.csv"
    source.head(0).to_csv(empty_sales, index=False)
    try:
        module.run(empty_sales, base / "project-1-empty-output")
    except ValueError as exc:
        assert "no sales records" in str(exc)
    else:
        raise AssertionError("Project 1 must reject an empty input extract")
    extreme_sales = pd.DataFrame([
        {"transaction_id": "T-HUGE-1", "order_date": "2026-04-01", "customer_id": "C-H1", "city": "Chennai", "product": "Desk", "quantity": 1e308, "unit_price": 1},
        {"transaction_id": "T-HUGE-2", "order_date": "2026-04-02", "customer_id": "C-H2", "city": "Chennai", "product": "Chair", "quantity": 1e308, "unit_price": 1},
    ])
    extreme_path = base / "project-1-aggregate-overflow.csv"
    extreme_sales.to_csv(extreme_path, index=False)
    try:
        module.run(extreme_path, base / "project-1-aggregate-overflow")
    except ValueError as exc:
        assert "revenue totals overflowed" in str(exc)
    else:
        raise AssertionError("Project 1 must reject an overflowing monthly revenue total")
    print("[PASS] Project 1: cleaning, duplicate/invalid-row audit, decision log and monthly reconciliation")


def verify_project_2(base: Path) -> None:
    folder = PROJECTS / "02-retail-performance"
    module = load_module("retail_performance_project", folder / "src" / "analyze_retail.py")
    result = module.analyze(folder / "data" / "retail_orders.csv", base / "project-2")
    assert len(result["orders"]) == 20
    assert result["orders"]["order_line_id"].is_unique
    assert result["orders"]["order_id"].duplicated().any(), "The sample must demonstrate multiple lines in one order"
    invalid_retail = pd.read_csv(folder / "data" / "retail_orders.csv")
    invalid_retail.loc[0, "customer_id"] = " "
    invalid_retail_path = base / "project-2-missing-customer.csv"
    invalid_retail.to_csv(invalid_retail_path, index=False)
    try:
        module.analyze(invalid_retail_path, base / "project-2-invalid")
    except ValueError as exc:
        assert "customer_id" in str(exc)
    else:
        raise AssertionError("Project 2 must reject blank customer identifiers")
    nonfinite_retail = pd.read_csv(folder / "data" / "retail_orders.csv")
    nonfinite_retail["unit_price"] = nonfinite_retail["unit_price"].astype(float)
    nonfinite_retail.loc[0, "unit_price"] = float("inf")
    nonfinite_path = base / "project-2-infinite-price.csv"
    nonfinite_retail.to_csv(nonfinite_path, index=False)
    try:
        module.analyze(nonfinite_path, base / "project-2-nonfinite")
    except ValueError as exc:
        assert "finite" in str(exc)
    else:
        raise AssertionError("Project 2 must reject non-finite monetary values")
    overflow_retail = pd.read_csv(folder / "data" / "retail_orders.csv").head(2).copy()
    overflow_retail["quantity"] = [1e308, 1e308]
    overflow_retail["unit_price"] = [1.0, 1.0]
    overflow_retail["discount_pct"] = [0.0, 0.0]
    overflow_path = base / "project-2-overflow.csv"
    overflow_retail.to_csv(overflow_path, index=False)
    try:
        module.analyze(overflow_path, base / "project-2-overflow")
    except ValueError as exc:
        assert "aggregation overflowed" in str(exc).lower()
    else:
        raise AssertionError("Project 2 must reject an overflowing segment/month aggregate")
    empty_retail = base / "project-2-empty.csv"
    pd.DataFrame(columns=pd.read_csv(folder / "data" / "retail_orders.csv").columns).to_csv(empty_retail, index=False)
    try:
        module.analyze(empty_retail, base / "project-2-empty-output")
    except ValueError as exc:
        assert "no order lines" in str(exc)
    else:
        raise AssertionError("Project 2 must reject an empty retail extract")
    extreme_retail = pd.read_csv(folder / "data" / "retail_orders.csv").head(2).copy()
    extreme_retail["order_line_id"] = ["HUGE-L1", "HUGE-L2"]
    extreme_retail["order_id"] = ["HUGE-O1", "HUGE-O2"]
    extreme_retail["customer_id"] = ["HUGE-C1", "HUGE-C2"]
    extreme_retail["quantity"] = [1e308, 1e308]
    extreme_retail["unit_price"] = [1.0, 1.0]
    extreme_retail["discount_pct"] = [0.0, 0.0]
    extreme_path = base / "project-2-aggregate-overflow.csv"
    extreme_retail.to_csv(extreme_path, index=False)
    try:
        module.analyze(extreme_path, base / "project-2-aggregate-overflow")
    except ValueError as exc:
        assert "aggregation overflowed" in str(exc)
    else:
        raise AssertionError("Project 2 must reject an overflowing segment aggregation")
    assert abs(float(result["monthly"]["net_revenue"].sum()) - result["net_revenue"]) < 0.01
    assert len(result["region"]) == 4 and len(result["category"]) == 3
    for name in ("monthly_performance.csv", "region_performance.csv", "category_performance.csv", "customer_performance.csv", "monthly_net_revenue.png", "region_performance.png", "category_performance.png", "insights.md"):
        assert (base / "project-2" / name).is_file(), name
    print("[PASS] Project 2: net revenue, segment summaries, three charts and business insight brief")


def verify_project_3(base: Path) -> None:
    folder = PROJECTS / "03-funnel-cohort-analysis"
    module = load_module("funnel_cohort_project", folder / "src" / "analyze_funnel_cohorts.py")
    result = module.analyze_events(folder / "data" / "customer_events.csv", base / "project-3")
    funnel = result["funnel"]
    counts = funnel["unique_customers"].tolist()
    assert counts == [12, 10, 6, 4], counts
    assert result["funnel"]["conversion_from_previous"].dropna().between(0, 1).all()

    # A customer only advances when each event occurs in sequence. The probe
    # includes a signup before visit, an early trial followed by a valid later
    # trial, and a purchase timestamp that precedes the trial.
    probe = pd.DataFrame([
        ("A1", "A", "visit", "2026-01-01 09:00"),
        ("A2", "A", "signup", "2026-01-01 10:00"),
        ("A3", "A", "trial_start", "2026-01-01 11:00"),
        ("A4", "A", "purchase", "2026-01-01 12:00"),
        ("B1", "B", "signup", "2026-01-01 08:00"),
        ("B2", "B", "visit", "2026-01-01 09:00"),
        ("B3", "B", "trial_start", "2026-01-01 10:00"),
        ("B4", "B", "purchase", "2026-01-01 11:00"),
        ("C1", "C", "visit", "2026-01-01 08:00"),
        ("C2", "C", "signup", "2026-01-01 09:00"),
        ("C3", "C", "trial_start", "2026-01-01 08:30"),
        ("C4", "C", "trial_start", "2026-01-01 10:00"),
        ("C5", "C", "purchase", "2026-01-01 11:00"),
        ("D1", "D", "visit", "2026-01-01 08:00"),
        ("D2", "D", "signup", "2026-01-01 09:00"),
        ("D3", "D", "trial_start", "2026-01-01 10:00"),
        ("D4", "D", "purchase", "2026-01-01 09:30"),
        ("E1", "E", "signup", "2026-01-01 08:00"),
        ("E2", "E", "trial_start", "2026-01-01 09:00"),
    ], columns=["event_id", "customer_id", "event_type", "event_timestamp"])
    probe["event_timestamp"] = pd.to_datetime(probe["event_timestamp"])
    probe_funnel = module.build_sequential_funnel(probe)
    assert probe_funnel["unique_customers"].tolist() == [4, 3, 3, 2]
    invalid_events = pd.read_csv(folder / "data" / "customer_events.csv", dtype={"event_id": "string", "customer_id": "string"})
    invalid_events.loc[0, "customer_id"] = ""
    invalid_path = base / "project-3-missing-customer.csv"
    invalid_events.to_csv(invalid_path, index=False)
    try:
        module.analyze_events(invalid_path, base / "project-3-invalid")
    except ValueError as exc:
        assert "customer_id" in str(exc)
    else:
        raise AssertionError("Project 3 must reject events without a customer_id")
    assert list(result["retention"].index) == ["2026-01", "2026-02", "2026-03"]
    assert float(result["cohort_activity"]["retention_pct"].min()) >= 0
    for name in ("funnel_metrics.csv", "cohort_activity.csv", "cohort_retention.csv", "cohort_retention_heatmap.png", "metric_definitions.md"):
        assert (base / "project-3" / name).is_file(), name
    print("[PASS] Project 3: distinct-customer funnel, monthly cohorts, retention matrix and definitions")


def verify_project_4(base: Path) -> None:
    folder = PROJECTS / "04-automated-reporting-pipeline"
    module = load_module("automated_reporting_project", folder / "src" / "report_pipeline.py")
    result = module.run_pipeline(folder / "fixtures" / "api_response.json", folder / "monthly", base / "project-4")
    rerun = module.run_pipeline(folder / "fixtures" / "api_response.json", folder / "monthly", base / "project-4")
    assert rerun["paid_revenue"] == result["paid_revenue"] and len(rerun["all_rows"]) == len(result["all_rows"])
    original_connection = sqlite3.connect(result["db_path"])
    try:
        original_rows = original_connection.execute("SELECT order_id, amount FROM orders ORDER BY order_id").fetchall()
    finally:
        original_connection.close()
    original_connect = module.sqlite3.connect
    def connection_that_rejects_order_insert(*args, **kwargs):
        connection = original_connect(*args, **kwargs)
        def authorizer(action, arg1, arg2, database_name, trigger_name):
            if action == sqlite3.SQLITE_INSERT and arg1 == "orders":
                return sqlite3.SQLITE_DENY
            return sqlite3.SQLITE_OK
        connection.set_authorizer(authorizer)
        return connection
    module.sqlite3.connect = connection_that_rejects_order_insert
    try:
        try:
            module.run_pipeline(folder / "fixtures" / "api_response.json", folder / "monthly", base / "project-4")
        except sqlite3.DatabaseError:
            pass
        else:
            raise AssertionError("Injected SQLite insert failure must abort the replacement load")
    finally:
        module.sqlite3.connect = original_connect
    preserved_connection = sqlite3.connect(result["db_path"])
    try:
        preserved_rows = preserved_connection.execute("SELECT order_id, amount FROM orders ORDER BY order_id").fetchall()
        assert preserved_rows == original_rows, "A failed replacement must preserve the previously committed orders table"
    finally:
        preserved_connection.close()
    assert len(result["api_records"]) == 3 and len(result["monthly_records"]) == 5
    assert len(result["all_rows"]) == 8 and len(result["paid"]) == 6
    assert abs(result["paid_revenue"] - 10200.0) < 0.01
    assert result["db_path"].is_file() and result["report_path"].is_file()
    for name in ("source_snapshot.json", "paid_order_extract.csv", "monthly_paid_summary.csv", "run.log"):
        assert (base / "project-4" / name).is_file(), name
    workbook = load_workbook(result["report_path"], data_only=False)
    try:
        assert workbook.sheetnames == ["Summary", "Details"]
        assert workbook["Summary"]["C5"].value == "=SUM(C2:C4)"
        assert workbook["Details"].max_row == 7
        assert workbook["Details"].freeze_panes == "A2"
    finally:
        workbook.close()
    assert "Reconciliation: PASS" in result["run_log"]
    valid_record = {
        "order_id": "PROBE-001", "month": "2026-04", "customer_id": "C-PROBE",
        "status": "paid", "amount": 12.5,
    }
    normalized = module.validate_order_records([valid_record], "probe")
    assert normalized[0]["amount"] == 12.5 and normalized[0]["status"] == "paid"
    invalid_cases = [
        ({**valid_record, "month": "2026-13"}, "invalid month"),
        ({**valid_record, "status": "unknown"}, "unsupported status"),
        ({**valid_record, "amount": float("inf")}, "finite and non-negative"),
        ({**valid_record, "amount": float("nan")}, "finite and non-negative"),
        ({**valid_record, "customer_id": " "}, "blank customer_id"),
        ({**valid_record, "order_id": None}, "missing order_id"),
    ]
    for invalid_record, expected_message in invalid_cases:
        try:
            module.validate_order_records([invalid_record], "probe")
        except ValueError as exc:
            assert expected_message in str(exc), str(exc)
        else:
            raise AssertionError(f"Project 4 must reject invalid record: {invalid_record}")

    # A valid batch with no paid orders should still produce a non-circular
    # workbook total instead of a reversed SUM range.
    pending_fixture = base / "pending-only.json"
    pending_fixture.write_text(json.dumps({"results": [{
        "order_id": "PENDING-001", "month": "2026-04", "customer_id": "C-PENDING",
        "status": "pending", "amount": 10.0,
    }]}), encoding="utf-8")
    empty_monthly = base / "empty-monthly"
    empty_monthly.mkdir()
    pending_result = module.run_pipeline(pending_fixture, empty_monthly, base / "project-4-pending")
    assert pending_result["paid"].empty
    overflow_fixture = base / "paid-overflow.json"
    overflow_fixture.write_text(json.dumps({"results": [
        {"order_id": "HUGE-101", "month": "2026-04", "customer_id": "C-H1", "status": "paid", "amount": 1e308},
        {"order_id": "HUGE-102", "month": "2026-04", "customer_id": "C-H2", "status": "paid", "amount": 1e308},
    ]}), encoding="utf-8")
    try:
        module.run_pipeline(overflow_fixture, empty_monthly, base / "project-4")
    except ValueError as exc:
        assert "previous report outputs were preserved" in str(exc)
    else:
        raise AssertionError("Project 4 must reject a non-finite paid revenue summary")
    preserved_after_overflow = sqlite3.connect(result["db_path"])
    try:
        rows_after_overflow = preserved_after_overflow.execute(
            "SELECT order_id, amount FROM orders ORDER BY order_id"
        ).fetchall()
        assert rows_after_overflow == original_rows
    finally:
        preserved_after_overflow.close()
    preserved_workbook = load_workbook(result["report_path"], data_only=False)
    try:
        assert preserved_workbook.sheetnames == ["Summary", "Details"]
        assert preserved_workbook["Summary"]["C5"].value == "=SUM(C2:C4)"
    finally:
        preserved_workbook.close()
    empty_workbook = load_workbook(pending_result["report_path"], data_only=False)
    try:
        assert empty_workbook["Summary"]["B2"].value == "=0"
        assert empty_workbook["Summary"]["C2"].value == "=0"
        assert empty_workbook["Details"].max_row == 1
    finally:
        empty_workbook.close()
    print("[PASS] Project 4: fixture ingestion, parameterised SQLite, pandas extract and Excel report")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="datatrail-programme1-projects-") as temporary:
        base = Path(temporary)
        verify_project_1(base)
        verify_project_2(base)
        verify_project_3(base)
        verify_project_4(base)
    print("[PASS] Programme 1 portfolio: four projects executed successfully in isolated temporary folders.")


if __name__ == "__main__":
    main()
