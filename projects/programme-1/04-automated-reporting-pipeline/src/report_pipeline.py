#!/usr/bin/env python3
"""Run an offline API-fixture → CSV → SQLite → pandas → Excel report pipeline."""

from __future__ import annotations

import argparse
import json
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def load_api_fixture(path: Path) -> list[dict]:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    records = payload.get("results")
    if not isinstance(records, list):
        raise ValueError("API fixture must contain a results list")
    return records


def load_monthly_csvs(folder: Path) -> list[dict]:
    folder = Path(folder)
    records = []
    required = {"order_id", "month", "customer_id", "status", "amount"}
    for path in sorted(folder.glob("*.csv")):
        frame = pd.read_csv(path)
        missing = required - set(frame.columns)
        if missing:
            raise ValueError(f"{path.name} is missing columns: {sorted(missing)}")
        if frame["order_id"].isna().any() or frame["order_id"].duplicated().any():
            raise ValueError(f"{path.name} contains a missing or duplicate order_id")
        frame["amount"] = pd.to_numeric(frame["amount"], errors="coerce")
        if frame["amount"].isna().any() or (frame["amount"] < 0).any():
            raise ValueError(f"{path.name} contains an invalid amount")
        if not frame["month"].astype(str).str.match(r"^\d{4}-\d{2}$").all():
            raise ValueError(f"{path.name} contains an invalid month")
        records.extend(frame.to_dict(orient="records"))
    return records


def run_pipeline(api_fixture: Path, monthly_folder: Path, output_dir: Path) -> dict:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    api_records = load_api_fixture(api_fixture)
    monthly_records = load_monthly_csvs(monthly_folder)
    all_rows = []
    for row in api_records:
        all_rows.append({**row, "source": "api_fixture"})
    for row in monthly_records:
        all_rows.append({**row, "source": "monthly_csv"})

    required = {"order_id", "month", "customer_id", "status", "amount", "source"}
    if any(required - set(row) for row in all_rows):
        raise ValueError("A source record is missing a required field")
    amounts = [float(row["amount"]) for row in all_rows]
    if any(value < 0 for value in amounts):
        raise ValueError("Order amounts cannot be negative")
    if len({row["order_id"] for row in all_rows}) != len(all_rows):
        raise ValueError("Order IDs must be unique across all sources")

    db_path = output_dir / "report_database.sqlite"
    connection = sqlite3.connect(db_path)
    try:
        connection.execute("PRAGMA foreign_keys = ON")
        with connection:
            connection.execute("""
                CREATE TABLE orders (
                    order_id TEXT PRIMARY KEY,
                    month TEXT NOT NULL,
                    customer_id TEXT NOT NULL,
                    status TEXT NOT NULL,
                    amount REAL NOT NULL CHECK(amount >= 0),
                    source TEXT NOT NULL
                )
            """)
            connection.executemany(
                "INSERT INTO orders (order_id, month, customer_id, status, amount, source) VALUES (?, ?, ?, ?, ?, ?)",
                [
                    (row["order_id"], row["month"], row["customer_id"], row["status"], float(row["amount"]), row["source"])
                    for row in all_rows
                ],
            )
        paid = pd.read_sql_query(
            "SELECT order_id, month, customer_id, amount, source FROM orders WHERE status = ? ORDER BY month, order_id",
            connection,
            params=("paid",),
        )
    finally:
        connection.close()

    summary = paid.groupby("month", as_index=False).agg(
        paid_orders=("order_id", "count"), paid_revenue=("amount", "sum")
    ).sort_values("month")
    if abs(float(summary["paid_revenue"].sum()) - float(paid["amount"].sum())) > 0.01:
        raise AssertionError("Monthly paid summary does not reconcile to database detail")

    snapshot = {"api_fixture_records": api_records, "monthly_csv_records": monthly_records}
    (output_dir / "source_snapshot.json").write_text(json.dumps(snapshot, indent=2, sort_keys=True), encoding="utf-8")
    paid.to_csv(output_dir / "paid_order_extract.csv", index=False)
    summary.to_csv(output_dir / "monthly_paid_summary.csv", index=False)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    detail_sheet = workbook.create_sheet("Details")
    summary_sheet.append(["Month", "Paid orders", "Paid revenue"])
    detail_sheet.append(["Order ID", "Month", "Customer ID", "Amount", "Source"])
    for row in paid.itertuples(index=False):
        detail_sheet.append([row.order_id, row.month, row.customer_id, float(row.amount), row.source])
    for row in summary.itertuples(index=False):
        summary_sheet.append([row.month, int(row.paid_orders), float(row.paid_revenue)])
    total_row = summary_sheet.max_row + 1
    summary_sheet.cell(total_row, 1, "Grand total")
    summary_sheet.cell(total_row, 2, f"=SUM(B2:B{total_row - 1})")
    summary_sheet.cell(total_row, 3, f"=SUM(C2:C{total_row - 1})")
    header_fill = PatternFill("solid", fgColor="173F35")
    for sheet in (summary_sheet, detail_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 36)
            sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
    for row in summary_sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = '#,##0.00'
    for row in detail_sheet.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = '#,##0.00'
    for cell in summary_sheet[total_row]:
        cell.font = Font(bold=True)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    report_path = output_dir / "monthly_report.xlsx"
    workbook.save(report_path)
    workbook.close()

    total_paid = float(paid["amount"].sum())
    run_log = (
        f"API fixture records: {len(api_records)}\n"
        f"Monthly CSV records: {len(monthly_records)}\n"
        f"Database rows: {len(all_rows)}\n"
        f"Paid orders: {len(paid)}\n"
        f"Paid revenue: {total_paid:.2f}\n"
        f"Reconciliation: PASS\n"
        f"Workbook: {report_path.name}\n"
    )
    (output_dir / "run.log").write_text(run_log, encoding="utf-8")
    return {
        "api_records": api_records, "monthly_records": monthly_records, "all_rows": all_rows,
        "paid": paid, "summary": summary, "report_path": report_path, "db_path": db_path,
        "paid_revenue": total_paid, "run_log": run_log,
    }


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--api-fixture", type=Path, default=project_root / "fixtures" / "api_response.json")
    parser.add_argument("--monthly-folder", type=Path, default=project_root / "monthly")
    parser.add_argument("--output-dir", type=Path, default=Path("output/project-4"))
    args = parser.parse_args()
    result = run_pipeline(args.api_fixture, args.monthly_folder, args.output_dir)
    print(f"Source rows: {len(result['all_rows'])}")
    print(f"Paid orders: {len(result['paid'])}")
    print(f"Paid revenue: {result['paid_revenue']:.2f}")
    print("Reconciliation: PASS")


if __name__ == "__main__":
    main()
